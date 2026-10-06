"""Carga de resultados del Congreso desde el dataset público del Ministerio del Interior."""
from __future__ import annotations
import io, re, ssl, urllib.request
from pathlib import Path
import json
from typing import Any
import certifi
import openpyxl

OFFICIAL_URL = "https://descargas.interior.gob.es/datasets/resultados_electorales/Elecciones-Congreso.xlsx"

def download_workbook(destination: str | Path) -> Path:
    dest = Path(destination)
    dest.parent.mkdir(parents=True, exist_ok=True)
    context = ssl.create_default_context(cafile=certifi.where())
    with urllib.request.urlopen(OFFICIAL_URL, context=context, timeout=60) as r:
        dest.write_bytes(r.read())
    return dest

def _norm(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())

def inspect_workbook(path: str | Path) -> dict[str, list[str]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = ws.iter_rows(min_row=1, max_row=min(ws.max_row, 5), values_only=True)
        headers = next(iter(rows), ())
        out[ws.title] = [str(x) for x in headers if x is not None]
    return out

def _find_col(headers, *names):
    normalized = {_norm(h): i for i, h in enumerate(headers)}
    for name in names:
        n = _norm(name)
        if n in normalized:
            return normalized[n]
    for i, h in enumerate(headers):
        n = _norm(h)
        if any(x in n for x in map(_norm, names)):
            return i
    return None

def load_rows(path: str | Path) -> list[dict[str, Any]]:
    """Return normalized rows when the workbook exposes tabular province/candidature data.

    The loader intentionally does not guess party families: labels remain source-native.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows_out = []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        headers = list(next(it, ()))
        if not headers:
            continue
        pcol = _find_col(headers, "Provincia", "Circunscripción", "Province")
        partycol = _find_col(headers, "Candidatura", "Candidaturas", "Partido", "Siglas")
        votescol = _find_col(headers, "Votos", "Votos candidatura", "Candidature votes")
        if pcol is None or partycol is None or votescol is None:
            continue
        for row in it:
            if len(row) <= max(pcol, partycol, votescol):
                continue
            province, party, votes = row[pcol], row[partycol], row[votescol]
            if province in (None, "") or party in (None, ""):
                continue
            if isinstance(votes, (int, float)) and votes >= 0:
                rows_out.append({"sheet": ws.title, "province": str(province).strip(), "party": str(party).strip(), "votes": int(votes)})
    if not rows_out:
        raise RuntimeError("No se encontró una tabla provincial/candidatura reconocible; use inspect_workbook() para adaptar el esquema.")
    return rows_out

def write_json(rows: list[dict[str, Any]], destination: str | Path) -> Path:
    dest = Path(destination)
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    return dest
