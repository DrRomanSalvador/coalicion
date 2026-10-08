#!/usr/bin/env python3
"""Materialize the official unified Interior Congress workbook.

The source remains the Ministry of the Interior workbook. This script can
regenerate the complete normalized historical dataset and the 2023 canonical
matrix from that source without inventing territorial values.
"""
from __future__ import annotations
import argparse, csv, hashlib, json
from pathlib import Path
import openpyxl
from src.data import load_official_constituency_matrix

OFFICIAL_URL="https://descargas.interior.gob.es/datasets/resultados_electorales/Elecciones-Congreso.xlsx"

def _constituency_columns(headers):
    out={}
    for index,value in enumerate(headers):
        if not isinstance(value,str) or " - " not in value: continue
        prefix,name=value.split(" - ",1)
        if prefix.isdigit() and 1<=int(prefix)<=52: out[index]=(int(prefix),name.strip())
    codes=[code for code,_ in out.values()]
    names=[name for _,name in out.values()]
    if len(codes)!=52 or set(codes)!=set(range(1,53)):
        raise ValueError("El workbook oficial no contiene exactamente una columna por cada una de las 52 circunscripciones.")
    if len(set(names))!=52 or any(not name for name in names):
        raise ValueError("Nombres oficiales de circunscripción vacíos o duplicados.")
    return out

def normalize_workbook(path:Path,destination:Path)->dict:
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    records=0; elections=set()
    destination.parent.mkdir(parents=True,exist_ok=True)
    # Publish only a fully validated dataset; failures never replace canonical output.
    temporary=destination.with_name(destination.name+".tmp")
    try:
        with temporary.open("w",encoding="utf-8",newline="") as fh:
            writer=csv.writer(fh)
            writer.writerow(["election_date","election_code","election_type","metric","subject","constituency_code","constituency","value","national_total"])
            for ws in wb.worksheets:
                header=next(ws.iter_rows(min_row=4,max_row=4,values_only=True),())
                cols=_constituency_columns(list(header))
                for row in ws.iter_rows(min_row=5,values_only=True):
                    if not row or row[0] is None: continue
                    if not hasattr(row[0],"date"): raise ValueError("Fecha de elección inválida.")
                    election_date=row[0].date().isoformat(); elections.add(election_date)
                    description=str(row[3] or "").strip()
                    if description.startswith("Votos (") and description.endswith(")"): metric,subject="votes",description[7:-1]
                    elif description.startswith("Escaños (") and description.endswith(")"): metric,subject="seats",description[9:-1]
                    else: metric,subject=description.lower().replace(" ","_"),""
                    raw_code=row[1]
                    if isinstance(raw_code,bool) or not isinstance(raw_code,(int,float)) or not float(raw_code).is_integer() or raw_code<1:
                        raise ValueError(f"Código electoral inválido: {election_date}:{raw_code!r}")
                    raw_national=row[56]
                    if raw_national is not None and (isinstance(raw_national,bool) or not isinstance(raw_national,(int,float)) or not float(raw_national).is_integer() or raw_national<0):
                        raise ValueError(f"Total nacional inválido: {election_date}:{description}:{raw_national!r}")
                    national_total="" if raw_national is None else int(raw_national)
                    for index,(code,constituency) in cols.items():
                        value=row[index] if index<len(row) else None
                        if value is None: continue
                        if isinstance(value,bool) or not isinstance(value,(int,float)) or not float(value).is_integer() or value<0:
                            raise ValueError(f"Valor inválido: {election_date}:{constituency}:{description}")
                        writer.writerow([election_date,int(raw_code),str(row[2]),metric,subject,code,constituency,int(value),national_total]); records+=1
        if len(elections)!=16:
            raise ValueError(f"Se esperaban 16 elecciones históricas; recibidas {len(elections)}.")
        temporary.replace(destination)
        return {"records":records,"elections":sorted(elections)}
    finally:
        wb.close()
        temporary.unlink(missing_ok=True)

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--input",default="data/raw/Elecciones-Congreso.xlsx")
    parser.add_argument("--output-csv",default="data/official_interior_congreso_1977_2023.csv")
    parser.add_argument("--canonical-2023",default="artifacts/data/election_2023_canonical.json")
    parser.add_argument("--manifest",default="data/manifests/official_interior_congreso.json")
    args=parser.parse_args()
    root=Path(__file__).resolve().parents[1]; source=root/args.input
    if not source.is_file(): raise SystemExit(f"BLOCKED: no existe el workbook oficial: {source}")
    payload=source.read_bytes()
    if not payload: raise SystemExit("BLOCKED: workbook oficial vacío")
    sha256=hashlib.sha256(payload).hexdigest()
    # Validate the canonical 2023 matrix before publishing any derived CSV.
    matrix=load_official_constituency_matrix(source,"2023-07-23")
    if matrix["validation"]["status"]!="PASS": raise SystemExit("BLOCKED: matriz oficial 2023 inválida")
    normalized=normalize_workbook(source,root/args.output_csv)
    result={"schema":"ELECTION_2023_CONSTITUENCY_MATRIX_V2","election":2023,"type":"general","source_tier":"OFFICIAL_PRIMARY","source":{"provider":"Ministerio del Interior / Infoelectoral","url":OFFICIAL_URL,"sha256":sha256,"hash_scope":"workbook_bytes"},"data":{"constituencies":matrix["constituencies"]},"validation":matrix["validation"]}
    out=root/args.canonical_2023; out.parent.mkdir(parents=True,exist_ok=True)
    out_tmp=out.with_name(out.name+".tmp")
    out_tmp.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    out_tmp.replace(out)
    manifest={"schema":"OFFICIAL_INTERIOR_CONGRESS_DATASET_V1","status":"READY","provider":"Ministerio del Interior / Infoelectoral","source_url":OFFICIAL_URL,"input":str(args.input),"sha256":sha256,"hash_scope":"workbook_bytes","bytes":len(payload),"normalized_csv":str(args.output_csv),"canonical_2023":str(args.canonical_2023),"records":normalized["records"],"elections":normalized["elections"],"constituencies":52,"validated_2023":matrix["validation"]}
    mp=root/args.manifest; mp.parent.mkdir(parents=True,exist_ok=True)
    mp_tmp=mp.with_name(mp.name+".tmp")
    mp_tmp.write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    mp_tmp.replace(mp)
    print(json.dumps(manifest,ensure_ascii=False))
if __name__=="__main__": main()
