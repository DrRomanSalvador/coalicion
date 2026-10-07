#!/usr/bin/env python3
"""Materialize official Interior Congress historical results without R."""
from __future__ import annotations
import csv, hashlib, json, ssl, urllib.request, subprocess, tempfile, os
import certifi
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
URL="https://descargas.interior.gob.es/datasets/resultados_electorales/Elecciones-Congreso.xlsx"
PAGE="https://infoelectoral.interior.gob.es/es/elecciones-celebradas/area-de-descargas/index.html"
XLSX=ROOT/"data/raw/Elecciones-Congreso.xlsx"
OUT=ROOT/"data/resultados_oficiales_2004_2023.csv"
MAN=ROOT/"data/manifests/INTERIOR_ACQUISITION.json"
EXPECTED_XLSX_SHA256="dba3394f1812f338067231bce68acf56af1e13ddf8cfb709a814bcc46357ebc2"
DATES={"2004":"2004-03-14","2008":"2008-03-09","2011":"2011-11-20","2015":"2015-12-20","2016":"2016-06-26","2019A":"2019-04-28","2019N":"2019-11-10","2023J":"2023-07-23"}
FNMT_SERVER_ROOT_PEM="""-----BEGIN CERTIFICATE-----
MIICbjCCAfOgAwIBAgIQYvYybOXE42hcG2LdnC6dlTAKBggqhkjOPQQDAzB4MQsw
CQYDVQQGEwJFUzERMA8GA1UECgwIRk5NVC1SQ00xDjAMBgNVBAsMBUNlcmVzMRgw
FgYDVQRhDA9WQVRFUy1RMjgyNjAwNEoxLDAqBgNVBAMMI0FDIFJBSVogRk5NVC1S
Q00gU0VSVklET1JFUyBTRUdVUk9TMB4XDTE4MTIyMDA5MzczM1oXDTQzMTIyMDA5
MzczM1oweDELMAkGA1UEBhMCRVMxETAPBgNVBAoMCEZOTVQtUkNNMQ4wDAYDVQQL
DAVDZXJlczEYMBYGA1UEYQwPVkFURVMtUTI4MjYwMDRKMSwwKgYDVQQDDCNBQyBS
QUlaIEZOTVQtUkNNIFNFUlZJRE9SRVMgU0VHVVJPUzB2MBAGByqGSM49AgEGBSuB
BAAiA2IABPa6V1PIyqvfNkpSIeSX0oNnnvBlUdBeh8dHsVnyV0ebAAKTRBdp20LH
sbI6GA60XYyzZl2hNPk2LEnb80b8s0RpRBNm/dfF/a82Tc4DTQdxz69qBdKiQ1oK
Um8BA06Oi6NCMEAwDwYDVR0TAQH/BAUwAwEB/zAOBgNVHQ8BAf8EBAMCAQYwHQYD
VR0OBBYEFAG5L++/EYZg8k/QQW6rcx/n0m5JMAoGCCqGSM49BAMDA2kAMGYCMQCu
SuMrQMN0EfKVrRYj3k4MGuZdpSRea0R7/DjiT8ucRRcRTBQnJlU5dUoDzBOQn5IC
MQD6SmxgiHPz7riYYqnOK8LZiqZwMR2vsJRM60/G49HzYqc8/5MuB1xJAWdpEgJy
v+c=
-----END CERTIFICATE-----
"""

def main():
    XLSX.parent.mkdir(parents=True,exist_ok=True)
    MAN.parent.mkdir(parents=True,exist_ok=True)
    data=None
    ca_candidates=[p for p in ("/etc/ssl/certs/ca-certificates.crt", certifi.where()) if os.path.exists(p)]
    if not ca_candidates:
        raise RuntimeError("FAIL-CLOSED: no trusted CA bundle available")
    ca=ca_candidates[0]
    bundle_tmp=Path(tempfile.mkstemp(suffix=".pem")[1])
    try:
        der=ssl.PEM_cert_to_DER_cert(FNMT_SERVER_ROOT_PEM)
        if hashlib.sha256(der).hexdigest().upper()!="554153B13D2CF9DDB753BFBE1A4E0AE08D0AA4187058FE60A2B862B2E4B87BCB":
            raise RuntimeError("FAIL-CLOSED: embedded FNMT server-root fingerprint mismatch")
        bundle_tmp.write_bytes(Path(ca).read_bytes()+b"\n"+FNMT_SERVER_ROOT_PEM.encode("ascii"))
        ca=bundle_tmp.as_posix()
    except Exception:
        bundle_tmp.unlink(missing_ok=True)
        raise
    with tempfile.NamedTemporaryFile(suffix=".xlsx",delete=False) as tmp:
        tmp_path=Path(tmp.name)
    curl_base=["curl","-fL","--http1.1","--retry","5","--retry-all-errors","--retry-delay","3",
               "--connect-timeout","30","--max-time","900","--compressed",
               "-A","coalicion-primary-materializer/1.0"]
    try:
        for cmd in (curl_base+["--cacert",ca,"-o",str(tmp_path),URL],
                    curl_base+["--cacert",certifi.where(),"-o",str(tmp_path),URL]):
            cp=subprocess.run(cmd,check=False,capture_output=True,text=True)
            if cp.returncode==0 and tmp_path.exists() and tmp_path.stat().st_size>10000:
                data=tmp_path.read_bytes()
                break
            tmp_path.unlink(missing_ok=True)
    finally:
        tmp_path.unlink(missing_ok=True)
        bundle_tmp.unlink(missing_ok=True)
    if data is None:
        # Last-resort transport only: TLS verification is unavailable on this
        # runner, so accept the official workbook exclusively when its bytes
        # match the repository's pre-established source hash. Never accept a
        # different payload.
        with tempfile.NamedTemporaryFile(suffix=".xlsx",delete=False) as tmp:
            tmp_path=Path(tmp.name)
        try:
            cmd=curl_base+["-k","-o",str(tmp_path),URL]
            cp=subprocess.run(cmd,check=False,capture_output=True,text=True)
            if cp.returncode==0 and tmp_path.exists() and tmp_path.stat().st_size>10000:
                candidate=tmp_path.read_bytes()
                digest=hashlib.sha256(candidate).hexdigest()
                if digest==EXPECTED_XLSX_SHA256:
                    data=candidate
                else:
                    raise RuntimeError("FAIL-CLOSED: official Interior payload hash mismatch")
        finally:
            tmp_path.unlink(missing_ok=True)
    if data is None:
        raise RuntimeError("FAIL-CLOSED: official Interior acquisition failed through trusted TLS and pinned-hash fallback")
    if len(data)<=10000 or data[:4]!=b"PK\x03\x04":
        raise RuntimeError("FAIL-CLOSED: official Interior response is not a valid XLSX")
    XLSX.write_bytes(data)
    sha=hashlib.sha256(data).hexdigest()
    wb=load_workbook(XLSX,read_only=True,data_only=True)
    if "Congreso" not in wb.sheetnames: raise RuntimeError("Missing Congreso sheet")
    ws=wb["Congreso"]
    rows=ws.iter_rows(values_only=True)
    headers=next(rows)
    # Official workbook has a title/header area; locate the real header row deterministically.
    header_rows=[headers]
    while headers and not any(str(v).strip().lower()=="fecha" for v in headers if v is not None):
        headers=next(rows)
    if not headers: raise RuntimeError("Fecha header not found")
    headers=list(headers)
    def col(*names):
        for i,v in enumerate(headers):
            if str(v).strip().lower() in {n.lower() for n in names}: return i
        return -1
    date_i=col("Fecha"); desc_i=col("Descripción","Descripcion"); type_i=col("Tipo Elección","Tipo Eleccion")
    if min(date_i,desc_i,type_i)<0: raise RuntimeError("Missing required columns")
    province_cols=list(range(4,56))
    if len(province_cols)!=52: raise RuntimeError("Expected 52 constituency columns")
    provinces=[str(headers[i]).strip() for i in province_cols]
    if any(not p or p.lower()=="nan" for p in provinces) or len(set(provinces))!=52: raise RuntimeError("Invalid constituency columns")
    out=[]
    for row in rows:
        if len(row)<=max(province_cols+[date_i,desc_i,type_i]): continue
        d=row[date_i]
        ds=d.strftime("%Y-%m-%d") if hasattr(d,"strftime") else str(d)[:10]
        election=next((e for e,x in DATES.items() if x==ds),None)
        if election is None or str(row[type_i]).strip().lower()!="congreso": continue
        desc=str(row[desc_i]).strip()
        import re
        m=re.match(r"^(Votos|Escaños|Escanos|Diputados)\s*\((.*)\)$",desc,re.I)
        if not m: continue
        metric="votos" if m.group(1).lower()=="votos" else "escaños"
        party=m.group(2).strip()
        if not party: continue
        for i,p in zip(province_cols,provinces):
            v=row[i]
            if isinstance(v,(int,float)) and v>=0:
                out.append((election,DATES[election],p,party,metric,int(v)))
    if not out: raise RuntimeError("No official rows recognized")
    keys={}
    for e,d,p,party,metric,v in out:
        keys.setdefault((e,d,p,party),{})[metric]=v
    if any(set(v)!= {"votos","escaños"} for v in keys.values()): raise RuntimeError("Missing/duplicate Votos/Escaños")
    result=[]
    for (e,d,p,party),v in sorted(keys.items()):
        result.append([e,d,p,party,v["votos"],v["escaños"],"Ministerio del Interior / portal oficial de datos abiertos","PRIMARY_OFFICIAL"])
    elections=set(r[0] for r in result)
    if elections!=set(DATES): raise RuntimeError(f"Missing elections: {sorted(set(DATES)-elections)}")
    seat_totals={e:sum(r[5] for r in result if r[0]==e) for e in DATES}
    constituency_totals={e:len({r[2] for r in result if r[0]==e}) for e in DATES}
    bad_seats={e:n for e,n in seat_totals.items() if n!=350}
    bad_const={e:n for e,n in constituency_totals.items() if n!=52}
    if bad_seats:
        raise RuntimeError(f"Historical seat totals are not 350: {bad_seats}")
    if bad_const:
        raise RuntimeError(f"Historical constituency totals are not 52: {bad_const}")
    with OUT.open("w",encoding="utf-8",newline="") as f:
        w=csv.writer(f); w.writerow(["election","fecha_eleccion","circunscripcion","partido","votos","escaños","fuente","nivel_fuente"]); w.writerows(result)
    manifest={"schema":"INTERIOR_OFFICIAL_ACQUISITION_V4","status":"PASS","source_url":URL,"download_page":PAGE,"sha256":sha,"file_bytes":len(data),"elections":DATES,"n_rows":len(result),"n_constituencies_per_election":52,"seats_per_election":350,"source_tier":"PRIMARY_INTERIOR","source_kind":"OFFICIAL_INTERIOR_XLSX","official_results":{"rows":len(result),"elections":list(DATES.values()),"constituencies_per_election":52,"seats_per_election":350,"derived_csv_sha256":hashlib.sha256(OUT.read_bytes()).hexdigest()}}
    MAN.write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(manifest,ensure_ascii=False,indent=2))

if __name__=="__main__": main()
